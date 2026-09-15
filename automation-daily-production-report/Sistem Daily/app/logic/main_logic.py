import win32com.client as win32

import os
import json
import logging
import time
import warnings
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

# Ignore openpyxl user warnings
warnings.filterwarnings("ignore", category=UserWarning)

# Setup logger
logging.basicConfig(
    level=logging.INFO,
    format="%(message)s",
    handlers=[logging.StreamHandler()]
)

# Site yang memakai logic Coal Winning per-hari
DAY_BASED_COAL_WINNING = {"IMM"}

# def force_recalc_with_excel(path):
#     excel = win32.gencache.EnsureDispatch('Excel.Application')
#     excel.Visible = False
#     wb = excel.Workbooks.Open(path)
#     wb.Save()
#     wb.Close(SaveChanges=True)
#     excel.Quit()

def force_recalc_with_excel(path):
    excel = win32.gencache.EnsureDispatch('Excel.Application')
    excel.Visible = False
    excel.DisplayAlerts = False

    wb = excel.Workbooks.Open(
        path,
        UpdateLinks=1,   # ⬅️ PENTING
        ReadOnly=False
    )

    wb.RefreshAll()                # pivot / query / link
    excel.CalculateFullRebuild()   # hitung ulang SEMUA formula
    wb.Save()
    wb.Close(SaveChanges=True)

    excel.Quit()


# Function to log information in a formatted style
def log_info(status, key, section, mode, source, target, value=None):
    logging.info(
        f"{status} "
        f"{key:<5} | "
        f"{section:<15} | "
        f"{mode:<6} | " 
        f"{source:<7} → {target:<7} | "
        f"{str(value):<20}"
    )

# Convert Excel cell coordinate (e.g., 'B2') to row and column indices
def excel_cell_to_index(cell_str):
    col_letter, row = coordinate_from_string(cell_str)
    col = column_index_from_string(col_letter)
    return int(row), int(col)

# Handle sum operations like AK1111+AK1112
def handle_sum_operation(ws_src, source_cell):
    """
    Handle sum operations for cells like AK1111+AK1112
    Returns the sum of multiple cells or single cell value
    """
    if '+' in source_cell:
        cells = source_cell.split('+')
        total_value = 0
        for cell in cells:
            cell = cell.strip()
            try:
                row, col = excel_cell_to_index(cell)
                cell_value = ws_src.cell(row=row, column=col).value or 0
                numeric_value = float(cell_value) if cell_value is not None else 0
                total_value += numeric_value
                logging.info(f"    📊 Sum operation: {cell} = {numeric_value}")
            except Exception as e:
                logging.warning(f"    ⚠️ Error reading cell {cell}: {e}")
                continue
        logging.info(f"    ✅ Total sum: {total_value}")
        return total_value
    else:
        try:
            src_row, src_col = excel_cell_to_index(source_cell)
            value = ws_src.cell(row=src_row, column=src_col).value
            if value is None or value == 0:
                return 0.0001
            try:
                return float(value)
            except:
                return value
        except Exception as e:
            logging.warning(f"    ⚠️ Error reading single cell {source_cell}: {e}")
            return 0.1

def get_coal_winning_daily_and_mtd_simple(ws, this_day: int):
    """
    DAILY:
      - cari 'sum of production' (kolom B) kemunculan ke-3
      - turun 2 baris -> target_row
      - ambil kolom: daily_col = 1 + this_day   (sesuai permintaan: mapping this date -1)
    MTD:
      - kolom di sebelah kanan daily -> mtd_col = daily_col + 1
    Return: (daily_value, mtd_value, target_row, daily_col, mtd_col)
    """
    if not (1 <= this_day <= 31):
        raise ValueError("This_Day harus 1..31")

    # cari kemunculan ke-3 'sum of production' di kolom B
    hit = 0
    third_row = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value  # kolom B
        s = (str(v).strip().lower() if v is not None else "")
        if s == "sum of production":
            hit += 1
            if hit == 3:
                third_row = r
                break

    if third_row is None:
        raise RuntimeError("Teks 'sum of production' tidak muncul 3 kali di kolom B.")

    target_row = third_row + 2
    if target_row > ws.max_row:
        raise RuntimeError("Baris target melewati akhir sheet.")

    # === offset kolom sesuai permintaan: mapping this date -1 ===
    daily_col = 1 + this_day      # (sebelumnya: 2 + this_day -> C utk day=1; sekarang dikurangi 1)
    mtd_col   = daily_col + 1     # sebelah kanannya daily

    # baca nilai
    daily_value = ws.cell(row=target_row, column=daily_col).value
    mtd_value   = ws.cell(row=target_row, column=mtd_col).value

    return daily_value, mtd_value, target_row, daily_col, mtd_col

def _to_float_or_zero(v):
    try:
        return float(v)
    except Exception:
        return 0.0

    
def get_rom_imm_value(ws): #Fungsi utama buat ngambil angka “ROM IMM” dari worksheet ws sesuai aturan West/East.
    """
    ROM IMM (revisi):
    - Temukan 'ROM Inv.' di kolom B
    - Di bawahnya: cari 'West' lalu 'East' (East setelah West; jika West tak ada, cari East setelah 'ROM Inv.')
    - Nilai ambil dari kolom AK
    - Logika:
        - Jika hanya West yang ada nilainya -> hasil = West
        - Jika hanya East yang ada nilainya -> hasil = East
        - Jika keduanya ada -> hasil = West + East (tanda +/- ikut)
        - Jika keduanya tak ada ATAU keduanya tak bernilai -> hasil = 0
    """
    def _to_float_or_none(v): #membersihkan nilai mentah jadi float. Kalau gak bisa, ya balikin None
        if v is None: #Kalau v = None 
            return None #langsung balikin None
        if isinstance(v, str):#Kalau v itu str (string)
            if v.strip() == "": #kalau setelah strip() (buang spasi kiri/kanan) jadi kosong ""
                return None #balikin None
            v = v.replace(",", "") #Kalau tidak kosong/terisi maka semua koma , dihapus contoh "1,234.5" jadi "1234.5" karna kalau tidak eror float() pyhton ga terima ,
        try: #Coba convert ke float dengan float(v)
            return float(v) #Kalau berhasil maka balikin nilai float-nya.
        except Exception:
            return None #Kalau gagal (ada error) maka balikin None.
        

    #nentuin nomor kolom Excel yang mau dipakai
    COL_B = 2 #kolom B (karena A=1, B=2, C=3, …).
    COL_AK = column_index_from_string("AK") 
    """
    fungsi openpyxl yang ngubah huruf kolom jadi angka.
    "AK" dikonversi jadi 37 (A=1… Z=26, AA=27, … AJ=36, AK=37).
    """

    # 1) cari 'ROM Inv.' di kolom B
    rominv_row = None #definisi variabel yang nilai awalnya kosong supaya nanti bisa diubah kalau ditemukan nilai yang benar.
    for r in range(1, ws.max_row + 1): #Loop dari baris 1 sampai baris terakhir di sheet (ws.max_row).
        s = ws.cell(row=r, column=COL_B).value #Ambil isi sel kolom B (karena COL_B = 2) di baris ke-r.
        s = (str(s).strip().lower() if s is not None else "")
        """
        Normalisasi nilainya:
        Kalau sel-nya None → jadi string kosong "".
        Kalau ada isinya → ubah ke str, buang spasi kiri/kanan (strip()), dan kecilkan huruf (lower()).
        """
        if s == "rom inv.": #Cek harus sama persis dengan "rom inv."
            rominv_row = r #Simpan nomor baris yang ketemu.
            break # lalu hentikan loop.
    if rominv_row is None: #Kalau dari awal sampai akhir nggak ada yang cocok
        raise RuntimeError("ROM Inv. tidak ditemukan di kolom B")
        #lempar RuntimeError("ROM Inv. tidak ditemukan di kolom B").

    # 2) cari 'West' (opsional)
    west_row = None #Inisialisasi: belum tahu posisinya, jadi diset kosong (None).
    for r in range(rominv_row + 1, ws.max_row + 1): #Loop dari baris setelah rominv_row (yaitu rominv_row + 1) sampai baris terakhir sheet.
        s = ws.cell(row=r, column=COL_B).value #Ambil isi sel pada kolom B di baris ke-r.
        s = (str(s).strip().lower() if s is not None else "")
        """
        Normalisasi nilai:
        Kalau sel kosong (None) → jadikan "".
        Kalau ada isinya → ubah ke string, buang spasi kiri/kanan (strip()), kecilkan huruf (lower()), biar perbandingan case-insensitive.
        """
        if s == "west": #Cek apakah teksnya persis “west” (setelah dinormalisasi).
            west_row = r #Kalau cocok, simpan nomor barisnya ke west_row.
            break #lalu hentikan loop.

    # 3) cari 'East' (setelah West kalau ada; jika West tak ada, setelah ROM Inv.)
    start_for_east = west_row if west_row is not None else rominv_row
    """
    Kalau west_row ada → mulai cari setelah West.
    Kalau west_row nggak ada → mulai cari setelah ROM Inv.
    """
    east_row = None #Inisialisasi: belum tahu baris “East”.
    for r in range(start_for_east + 1, ws.max_row + 1): #Loop dari baris setelah start_for_east sampai akhir sheet.
        s = ws.cell(row=r, column=COL_B).value #ambil isi kolom B baris r.
        s = (str(s).strip().lower() if s is not None else "") #Normalisasi (hapus spasi pinggir, jadikan huruf kecil, sel kosong jadi "").
        if s == "east": #Kalau teksnya pas “east”.
            east_row = r #simpan east_row = r.
            break #berhenti.

        #Kalau sampai habis loop nggak nemu, east_row tetap None (artinya “East” tidak ada).

    # 4) baca nilai AK
    """
    west_row / east_row: nomor baris tempat tulisan “West”/“East” ditemukan sebelumnya. Bisa berupa angka (mis. 12) atau None kalau nggak ketemu.
    COL_AK: nomor kolom untuk “AK” (mis. 37).
    ws.cell(row=..., column=COL_AK).value: baca isi sel di baris itu, kolom AK.
    _to_float_or_none(...): fungsi yang kamu punya untuk:
        hapus koma ribuan,
        coba ubah ke float,
        kalau kosong/invalid → None.
    Operator kondisionalnya A if kondisi else B:
        Kalau west_row ada nilainya (truthy) → ambil sel (west_row, AK) → parse → simpan ke west_val.
        Kalau tidak (west_row None) → west_val = None.
    Sama persis untuk east_val.
    """
    west_val = _to_float_or_none(ws.cell(row=west_row, column=COL_AK).value) if west_row else None
    east_val = _to_float_or_none(ws.cell(row=east_row, column=COL_AK).value) if east_row else None

        # 5) aturan hasil (kondisi gabungan) 
    #menggabungkan nilai West & East jadi satu angka result pakai aturan prioritas.
    if (west_row is None and east_row is None) or (west_val is None and east_val is None):
        result = 0.0
        """
        Kalau dua baris nggak ada atau dua nilainya None → result = 0.0
        (artinya nggak ada data yang bisa dipakai)
        """
    elif west_val is None:
        result = east_val
        """
        Kalau west_val = None saja → result = east_val
        (hanya East yang punya angka)
        """
    elif east_val is None:
        result = west_val
        """
        Kalau east_val = None saja → result = west_val
        (hanya West yang punya angka)
        """
    else:
        result = west_val + east_val
        """
        Kalau dua-duanya ada angka → result = west_val + east_val
        (jumlahkan; kalau ada yang negatif, tandanya ikut nilai aslinya)
        """

    # --- NEW: kalau hasil akhir negatif, paksa jadi 0 ---
    if result is None:
        result = 0.0
    elif result < 0:
        logging.info(f"ℹ️ ROM IMM result negatif ({result}), di-set jadi 0")
        result = 0.0
        
    #nulis catatan ke log pada level INFO: nomor baris “West/East”, nilai AK yang sudah diparse, dan hasil akhirnya.    
    logging.info(
        f"🏷️ ROM IMM: row West={west_row} (AK={west_val}), "
        f"row East={east_row} (AK={east_val}), result={result}"
    )
    return result

wb_cache = {} # tempat nyimpen workbook yang sudah pernah dibuka

# Load workbook with caching support
"""
Saat satu proses/skrip berkali-kali akses worksheet dari file yang sama (mis. baca beberapa sheet atau hitung beberapa bagian). 
Tanpa cache, tiap kali bisa lambat karena file dibuka berulang.
"""
def get_workbook_cached(path, read_only=False, data_only=True):
    key = f"{path}|{read_only}|{data_only}" #key = gabungan path, read_only, data_only.
    if key not in wb_cache: # kalau belum ada di cache
        wb_cache[key] = load_workbook( # buka file sekali saja
            path, data_only=data_only, read_only=read_only)
    return wb_cache[key] #langsung return objek dari cache (tanpa buka file lagi).

# NEW: Separate ROM Stock processing function
"""
fungsi ini otomatis baca & tulis data “ROM Stock” dari banyak file Excel sesuai konfigurasi (config) lalu nulis hasilnya ke workbook tujuan (wb_dst). 
Ada aturan khusus untuk site “IMM” (pakai rumus West/East kolom AK), kalau gagal/atau site lain → fallback ke pembacaan standar (handle_sum_operation).
"""
def process_rom_stock_automatically(config, file_map, wb_dst):
    """
    Process ROM Stock data automatically without manual input
    ROM Stock data is extracted directly from production files and processed to Excel
    """
    logging.info("[INFO] 🔄 Starting automatic ROM Stock processing...") #nulis log awal.
    for main_key, sections in config.items(): #Loop semua entri di config
 
        # skip yang bukan struktur utama: "file", "date", "This_Month", "Pro_Outlook_Sheet_Name".
        if main_key in ["file", "date", "This_Month", "Pro_Outlook_Sheet_Name"]:
            continue

        if not isinstance(sections, dict):#Pastikan struktur benar: sections harus dict. Kalau bukan, tulis peringatan dan lanjut.
            logging.warning(f"[WARN] Skip {main_key}, karena bukan dict (isi: {type(sections).__name__})")
            continue

        for section_name, categories in sections.items():
            for category_name, modes in categories.items():
                # Only process ROM Stock category
                if category_name != "ROM Stock":
                    continue

                # ROM Stock uses the same file as the main key
                source_key = main_key
                source_path = file_map.get(source_key)

                if not source_path:
                    log_info("⚠️", source_key, "ROM_AUTO", "-", "-", "-", "Source file not found")
                    continue

                try:
                    wb_src = get_workbook_cached(source_path, read_only=True)
                except Exception as e:
                    log_info("❌", source_key, "ROM_AUTO", "-", "-", "-", f"Failed to open: {e}")
                    continue

                for mode, mappings in modes.items():
                    for map_item in mappings:
                        try:
                            source_sheet = map_item.get("source_sheet")
                            source_cell = map_item.get("source_cell")
                            target_sheet = map_item.get("target_sheet")
                            target_cell = map_item.get("target_cell")

                            if not all([source_sheet, source_cell, target_sheet, target_cell]):
                                log_info("⚠️", source_key, "ROM_AUTO", mode, source_cell or "?", target_cell or "?", "Incomplete mapping")
                                continue

                            if source_sheet not in wb_src.sheetnames:
                                log_info("❌", source_key, "ROM_AUTO", mode, source_cell, target_cell, f"Sheet '{source_sheet}' not found")
                                continue

                            ws_src = wb_src[source_sheet]

                            # === khusus ROM IMM override: gunakan aturan West/East kolom AK
                            value = None
                            if main_key == "IMM":
                                try:
                                    value = get_rom_imm_value(ws_src) #Coba hitung value = get_rom_imm_value(ws_src) (pakai aturan West/East kolom AK).
                                    ws_dst = wb_dst[target_sheet] if target_sheet in wb_dst.sheetnames else wb_dst.create_sheet(title=target_sheet)
                                    dst_row, dst_col = excel_cell_to_index(target_cell)
                                    ws_dst.cell(row=dst_row, column=dst_col, value=value)
                                    log_info("✅", source_key, "ROM_AUTO", mode, "ROM Inv. West/East@AK", target_cell, value)
                                    continue  # selesai mapping ini, lanjut ke item berikutnya
                                except Exception as e:
                                    logging.warning(f"⚠️ ROM IMM rule gagal: {e}. Fallback ke mapping/source_cell biasa.")

                            # === fallback umum (site selain IMM, atau rule IMM gagal)
                            value = handle_sum_operation(ws_src, source_cell)

                            # Write to destination
                            #Siapkan sheet tujuan di wb_dst (buat kalau belum ada).
                            ws_dst = wb_dst[target_sheet] if target_sheet in wb_dst.sheetnames else wb_dst.create_sheet(title=target_sheet)
                            dst_row, dst_col = excel_cell_to_index(target_cell)
                            ws_dst.cell(row=dst_row, column=dst_col, value=value)

                            log_info("✅", source_key, "ROM_AUTO", mode, source_cell, target_cell, value)
                        except Exception as e:
                            log_info("❌", source_key, "ROM_AUTO", mode, map_item.get("source_cell", "?"), map_item.get("target_cell", "?"), f"Error: {e}")

    logging.info("[INFO] ✅ ROM Stock automatic processing completed!") #tulis log “completed!”


# -----------------------------
# COPY-PASTE: Process Summary -> Output 'Daily Production Report'
# -----------------------------
def process_summary_to_output(config, file_map, wb_dst):
    logging.info("[INFO] 🔄 Starting Summary -> Output copy process...")

    summary_path = file_map.get("Summary")
    if not summary_path or not os.path.exists(summary_path):
        logging.warning("[WARN] Summary file not provided or not found; skipping Summary -> Output step.")
        return

    pro_outlook_sheet_name = config.get("Pro_Outlook_Sheet_Name", "Outlook")

    try:
        wb_summary = get_workbook_cached(summary_path, read_only=False)
    except Exception as e:
        logging.error(f"[ERROR] Failed to open Summary workbook: {e}")
        return

    if pro_outlook_sheet_name not in wb_summary.sheetnames:
        logging.error(f"[ERROR] Sheet '{pro_outlook_sheet_name}' not found in Summary workbook.")
        return

    ws_summary = wb_summary[pro_outlook_sheet_name]

    month_col_map = {
        "Jan": "D", "Feb": "E", "Mar": "F", "Apr": "G",
        "May": "H", "Jun": "I", "Jul": "J", "Aug": "K",
        "Sep": "L", "Oct": "M", "Nov": "N", "Dec": "O"
    }

    this_month = config.get("This_Month")
    if not this_month:
        logging.warning("[WARN] This_Month not found in config; skipping Summary -> Output step.")
        return

    month_col_letter = month_col_map.get(this_month)
    if not month_col_letter:
        logging.warning(f"[WARN] This_Month '{this_month}' is not in month map; skipping.")
        return
    month_col_idx = column_index_from_string(month_col_letter)

    company_label_map = {
        "IMM": "FC Ready to Haul Product (Kton)",
        "BEK": "FC Product (KTons)",
        "TCM": "FC Production TCM (Kton)",
        "JBG": "JBG FC Product (Kton)",
        "GPK": "FC Product (Kton)",
        "TIS": "FC Product (Kton)",
        "NPR": "FC Product (Kton)"
    }

    target_sheet_name = "Daily Production Report"
    ws_out = wb_dst[target_sheet_name] if target_sheet_name in wb_dst.sheetnames else wb_dst.create_sheet(title=target_sheet_name)

    def find_first_row_with_value(ws, col_idx, value, start_row=1):
        if ws is None:
            logging.error("[ERROR] Worksheet is None in find_first_row_with_value")
            return None

        max_rows = ws.max_row
        if max_rows is None or max_rows < 1:
            logging.error(f"[ERROR] Worksheet {ws.title} has invalid max_row={max_rows}")
            return None

        for r in range(start_row, max_rows + 1):
            cell_value = ws.cell(row=r, column=col_idx).value
            if str(cell_value).strip() == str(value).strip():
                return r
        return None

    col_site_idx = column_index_from_string("C")
    col_fc_production_idx = column_index_from_string("F")
    col_fc_plan_idx = column_index_from_string("M")
    col_fc_pct_idx = column_index_from_string("L")

    for company, summary_label in company_label_map.items():
        try:
            logging.info(f"[INFO] Processing company: {company}")
            logging.info(f"[DEBUG] ws_out={ws_out}, sheetnames={wb_dst.sheetnames}")
            logging.info(f"[DEBUG] ws_out.max_row={getattr(ws_out, 'max_row', None)}")

            summary_company_row = find_first_row_with_value(ws_summary, column_index_from_string("B"), company, start_row=1)
            if not summary_company_row:
                logging.warning(f"[WARN] Company '{company}' not found in column B of Summary; skipping.")
                continue

            search_row = None
            safe_max = ws_summary.max_row if getattr(ws_summary, "max_row", None) else 10000
            for r in range(summary_company_row + 1, safe_max + 1):
                val = ws_summary.cell(row=r, column=column_index_from_string("B")).value
                if val is None:
                    continue
                if str(val).strip() == str(summary_label).strip():
                    search_row = r
                    break

            if not search_row:
                logging.warning(f"[WARN] Metric '{summary_label}' not found under company '{company}' in Summary; skipping.")
                continue

            value_cell = ws_summary.cell(row=search_row, column=month_col_idx).value
            try:
                numeric_value = 0 if value_cell is None else float(value_cell)
                numeric_value *= 1000
            except Exception:
                numeric_value = value_cell

            logging.info(f"[INFO] Found value for {company} at Summary row {search_row}, col {month_col_letter}: {numeric_value}")

            out_row = find_first_row_with_value(ws_out, col_site_idx, company, start_row=1)
            if not out_row:
                logging.warning(f"[WARN] Company '{company}' not found in Output sheet column C; skipping writing FC Plan.")
                continue

            ws_out.cell(row=out_row, column=col_fc_plan_idx, value=numeric_value)
            log_info("✅", company, "Summary->Output", "-", f"{summary_label}@{month_col_letter}{search_row}", f"M{out_row}", numeric_value)

            excel_formula = f"=IFERROR(SUMIF(C:C,\"{company}\",F:F)/M{out_row},0)"
            ws_out.cell(row=out_row, column=col_fc_pct_idx, value=excel_formula)
            log_info("✅", company, "Summary->Output", "-", "Formula", f"L{out_row}", excel_formula)

        except Exception as e:
            logging.exception(f"[ERROR] Error processing company {company}: {e}")

    logging.info("[INFO] ✅ Summary -> Output copy process completed!")

# === MAIN FUNCTION CALLED FROM GUI ===
def run_main_logic(file_map: dict):
    start_time = time.time()

    # Load full configuration from file
    mapping_file = 'config/inputan.json'
    try:
        with open(mapping_file, 'r') as f:
            config = json.load(f)
    except Exception as e:
        logging.error(f"[ERROR] Failed to open config file: {e}")
        return

    # Note: file_map from GUI no longer contains ROM Stock entry
    config["file"] = file_map
    file_map = config["file"]

    output_file = file_map.get("Output")
    if not output_file or not os.path.exists(output_file):
        logging.error("[ERROR] Output file is invalid or not found.")
        return

    try:
        wb_dst = load_workbook(output_file)
    except Exception as e:
        logging.error(f"[ERROR] Failed to open output file: {e}")
        return
    
    # >>> Ambil this_day dari config (pakai kunci 'This_Day' atau fallback 'this_Day')
    this_day = None
    try:
        if "This_Day" in config:
            this_day = int(config["This_Day"])
        elif "this_Day" in config:
            this_day = int(config["this_Day"])
    except Exception:
        this_day = None

    # === Special Date Mapping ===
    if "date" in config:
        date_map = config["date"]
        main_key = date_map.get("main_key", "BEK")
        source_path = file_map.get(main_key)
        logging.info(f"[INFO] This_Day = {this_day}")

        if source_path:
            try:
                wb_src = get_workbook_cached(source_path, read_only=True)
                source_sheet = date_map.get("source_sheet")
                source_cell = date_map.get("source_cell")
                target_sheet = date_map.get("target_sheet")
                target_range = date_map.get("target_range")

                if all([source_sheet, source_cell, target_sheet, target_range]):
                    ws_src = wb_src[source_sheet]
                    src_row, src_col = excel_cell_to_index(source_cell)
                    value = ws_src.cell(row=src_row, column=src_col).value
                    if value is None or value == 0:
                        value = 0.0001
                    ws_dst = wb_dst[target_sheet] if target_sheet in wb_dst.sheetnames else wb_dst.create_sheet(title=target_sheet)
                    start_cell, end_cell = target_range.split("-")
                    start_row, start_col = excel_cell_to_index(start_cell)
                    end_row, end_col = excel_cell_to_index(end_cell)

                    for row in range(start_row, end_row + 1):
                        for col in range(start_col, end_col + 1):
                            ws_dst.cell(row=row, column=col, value=value)

                    log_info("✅", main_key, "DateMapping", "-", source_cell, target_range, value)
                else:
                    log_info("⚠️", main_key, "DateMapping", "-", source_cell or "?", target_range or "?", "Incomplete mapping")
            except Exception as e:
                log_info("❌", main_key, "DateMapping", "-", source_cell or "?", target_range or "?", f"Error: {e}")
        else:
            log_info("⚠️", main_key, "DateMapping", "-", "-", "-", "Source file not available")

    # === ROM Stock Automatic Processing (FIRST) ===
    process_rom_stock_automatically(config, file_map, wb_dst)

    # === Main Mapping Process (EXCLUDING ROM STOCK) ===
    for main_key, sections in config.items():
        if main_key in ["file", "date", "This_Month", "Pro_Outlook_Sheet_Name"]:
            continue

        if not isinstance(sections, dict):
            logging.warning(f"[WARN] Skip {main_key}, karena bukan dict (isi: {type(sections).__name__})")
            continue

        for section_name, categories in sections.items():
            for category_name, modes in categories.items():

                if category_name == "ROM Stock":
                    continue

                # Tentukan file sumber
                if main_key == "IMM" and category_name in ["FC Production", "Port Stock Yard"]:
                    source_key = "SUM" 
                elif main_key == "IMM" and category_name == "Coal Winning":
                    source_key = "IMM"
                else:
                    source_key = main_key

                source_path = file_map.get(source_key)
                if not source_path:
                    log_info("⚠️", source_key, section_name, "-", "-", "-", "Source file not found, skipped")
                    continue

                try:
                    wb_src = get_workbook_cached(source_path, read_only=True)
                except Exception as e:
                    log_info("❌", source_key, section_name, "-", "-", "-", f"Failed to open source file: {e}")
                    continue

                for mode, mappings in modes.items():
                    for map_item in mappings:
                        try:
                            source_sheet = map_item.get("source_sheet")
                            source_cell = map_item.get("source_cell")
                            target_sheet = map_item.get("target_sheet")
                            target_cell = map_item.get("target_cell")

                            if not all([source_sheet, source_cell, target_sheet, target_cell]):
                                log_info("⚠️", source_key, section_name, mode, source_cell or "?", target_cell or "?", "Incomplete mapping")
                                continue

                            if source_sheet not in wb_src.sheetnames:
                                log_info("❌", source_key, section_name, mode, source_cell, target_cell, f"Sheet '{source_sheet}' not found")
                                continue

                            ws_src = wb_src[source_sheet]

                           # === Pilih cara ambil value ===
                            value = None

                            if (
                                (main_key in DAY_BASED_COAL_WINNING)
                                and (category_name == "Coal Winning")
                                and isinstance(this_day, int)
                                and 1 <= this_day <= 31
                            ):
                                try:
                                    daily_val, mtd_val, r, c_daily, c_mtd = get_coal_winning_daily_and_mtd_simple(ws_src, this_day)
                                    mode_norm = str(mode).strip().lower()
                                    if "daily" in mode_norm:
                                        value = daily_val
                                        logging.info(f"⛏️ Coal Winning {main_key} DAILY (this_day={this_day}) -> row={r}, col={c_daily} => {value}")
                                    elif "mtd" in mode_norm:
                                        value = mtd_val
                                        logging.info(f"⛏️ Coal Winning {main_key} MTD (kanan daily) -> row={r}, col={c_mtd} => {value}")
                                    # selain 'daily'/'mtd' tetap fallback
                                except Exception as e:
                                    logging.warning(f"⚠️ Coal Winning simple-index gagal untuk {main_key} [{mode}]: {e}. Fallback ke mapping source_cell.")

                            # Fallback umum kalau bukan kasus by-day / error
                            if value is None:
                                value = handle_sum_operation(ws_src, source_cell)

                            ws_dst = wb_dst[target_sheet] if target_sheet in wb_dst.sheetnames else wb_dst.create_sheet(title=target_sheet)
                            dst_row, dst_col = excel_cell_to_index(target_cell)
                            ws_dst.cell(row=dst_row, column=dst_col, value=value)

                            log_info("✅", source_key, section_name, mode, source_cell, target_cell, value)
                        except Exception as e:
                            log_info("❌", source_key, section_name, mode, map_item.get("source_cell", "?"), map_item.get("target_cell", "?"), f"Error: {e}")

    # Tambahkan: proses copy dari Summary -> Output sebelum save
    try:
        process_summary_to_output(config, file_map, wb_dst)
    except Exception as e:
        logging.exception(f"[ERROR] process_summary_to_output failed: {e}")

    # Save the workbook
    wb_dst.save(output_file)
    force_recalc_with_excel(output_file)
    end_time = time.time()

    # Final summary
    logging.info(f"[INFO] ✅ All processing completed successfully!")
    logging.info(f"[INFO] 📁 File saved: {output_file}")
    logging.info(f"[INFO] ⏱️  Total processing time: {end_time - start_time:.2f} seconds")
    logging.info(f"[INFO] 🔄 ROM Stock processed automatically from production files")
    logging.info(f"[INFO] 📊 Draft Excel ready for Power BI Dashboard")
