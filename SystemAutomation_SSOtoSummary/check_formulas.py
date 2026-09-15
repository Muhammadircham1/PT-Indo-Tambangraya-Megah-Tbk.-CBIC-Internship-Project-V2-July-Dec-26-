import openpyxl
wb = openpyxl.load_workbook(r'd:\magang\ITMG-Internship-All-Project-2025\SystemAutomation_SSOtoSummary\data\10 August_Margin.xlsx')
ws = wb['ITM Summary']

for r in range(250, 300):
    val_b = ws.cell(row=r, column=2).value
    if val_b == 'May':
        # found the header of May
        start_row = r + 3
        print(f"May starts at row {start_row}")
        for current_row in range(start_row, start_row + 5):
            print(f"\nRow {current_row}:")
            for col in ['AOV', 'AOW', 'AOX', 'AOY', 'AOZ', 'APA', 'APB', 'APC', 'APD', 'APE', 'APF', 'APG', 'APH']:
                print(f"{col}: {ws[col + str(current_row)].value}")
        break
