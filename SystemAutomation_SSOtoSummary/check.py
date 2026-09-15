import openpyxl
wb = openpyxl.load_workbook(r'd:\magang\ITMG-Internship-All-Project-2025\SystemAutomation_SSOtoSummary\data\10 August_Margin.xlsx', data_only=True)
ws = wb['FC Quality Master']

for r in range(1, 1600):
    c1 = ws.cell(row=r, column=1).value
    c2 = ws.cell(row=r, column=2).value
    for c in [c1, c2]:
        if c and isinstance(c, str) and ('Jan' in c or 'Feb' in c or 'Oct' in c or 'Sep' in c or 'Nov' in c or 'Dec' in c):
            print('Row:', r, 'Value:', c)
        elif c and hasattr(c, 'month') and hasattr(c, 'year'):
            if c.month in [1, 2, 9, 10, 11, 12]:
                print('Row:', r, 'Date:', c)

