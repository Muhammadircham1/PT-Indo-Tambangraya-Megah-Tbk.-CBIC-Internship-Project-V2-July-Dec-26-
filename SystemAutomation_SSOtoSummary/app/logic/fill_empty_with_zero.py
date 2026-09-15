from openpyxl.utils import column_index_from_string

def fill_empty_range_with_zero(sheet, check_col="A", start_col="O", end_col="AY", start_row=2):
    """
    Isi cell kosong dengan 0 di sheet tertentu.
    Hanya berlaku jika kolom A pada baris tersebut ada isinya.
    Range target: kolom O sampai AY (default).
    """
    check_idx = column_index_from_string(check_col)
    start_idx = column_index_from_string(start_col)
    end_idx = column_index_from_string(end_col)

    for row in range(start_row, sheet.max_row + 1):
        if sheet.cell(row=row, column=check_idx).value not in (None, ""):
            for col in range(start_idx, end_idx + 1):
                cell = sheet.cell(row=row, column=col)
                if cell.value in (None, ""):
                    cell.value = 0