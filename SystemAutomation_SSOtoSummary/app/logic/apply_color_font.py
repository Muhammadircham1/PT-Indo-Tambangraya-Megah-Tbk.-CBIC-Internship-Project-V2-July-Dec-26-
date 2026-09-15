from openpyxl.styles import Font

def apply_status_font(sheet, status_col_idx=69, start_col=14, end_col=62):
    """
    Change the color font in the blending column (N to BJ) based on status (kolom BQ).
    - Status == "Plan" / "Loading" / "In Progress" → font blue Verdana size 9
    - Status == "Completed" → font black Verdana size 9
    """
    max_row = sheet.max_row

    for row in range(2, max_row + 1):  # skip header row (start from row 2)
        status_val = sheet.cell(row=row, column=status_col_idx).value
        if not status_val:
            continue

        # set the font color
        if str(status_val).strip().lower() in ["plan", "loading", "in progress"]:
            font_color = "0070C0"  # blue
        elif str(status_val).strip().lower() == "completed":
            font_color = "000000"  # black
        else:
            continue  # unknown status → leave as default

        # loop column N (14) until BJ (62)
        for col in range(start_col, end_col + 1):
            cell = sheet.cell(row=row, column=col)
            if cell.value is not None:
                cell.font = Font(name="Verdana", size=9, color=font_color)

    print("🎨 Font Color in column blending successfully applied (Plan/Loading/In Progress = blue, Completed = black)")