from openpyxl.styles import PatternFill, Color
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

def backup_fpg(ws_summary: Worksheet, plan_sheet_name: str, comp_sheet_name: str, start_row: int, end_row: int):
    """
    Backup FPG Actual Value (ANQ) to the sheets FPG_plan and FPG_comp
    based on Status (BQ).
    """
    wb = ws_summary.parent

    # === 1. Create temporary backup sheet if not exist ===
    def create_temp_sheet(wb, sheet_name):
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            headers = ["Month", "Company", "Vessel", "Buyer", "End User", "Value", "FillColor"]
            ws.append(headers)
        else:
            ws = wb[sheet_name]

        return ws

    ws_plan = create_temp_sheet(wb, plan_sheet_name)
    ws_comp = create_temp_sheet(wb, comp_sheet_name)

    # === 2. Get the Column Index ===
    try:
        ANQ_col = column_index_from_string("ANQ")
        BQ_col  = column_index_from_string("BQ")
    except ValueError:
        raise ValueError("ERROR: 'ANQ' or 'BQ' is not valid Excel Column Name")

    print(f"\n=== 📦 BACKUP FPG VALUE ({start_row} → {end_row}) ===")

    # === 3. Extract fill color ===
    def extract_fill_hex(cell):
        fill = cell.fill
        if not fill or not fill.fgColor:
                    return "FFFFFFFF"  # default white

        rgb = getattr (fill.fgColor, "rgb", None)
        if isinstance(rgb, str):
            return rgb  # direct RGB/ARGB

        if fill.fgColor.type == "theme":
            return f"THEME_{fill.fgColor.theme}_TINT_{fill.fgColor.tint}"

        if fill.fgColor.type == "indexed":
            return f"INDEXED_{fill.fgColor.indexed}"

        return "UNKNOWN"

    # === 4. Loop rows for Backup ===
    for r in range(start_row, end_row + 1):
        cell_anq = ws_summary.cell(row=r, column=ANQ_col)
        val = cell_anq.value

        # skip if empty or formula
        if val is None:
            continue
        if isinstance(val, str) and val.startswith("="):
            continue  # formula, skip

        # Get KEY 5 Column C-G
        key = [
            ws_summary.cell(row=r, column=3).value,  # Month
            ws_summary.cell(row=r, column=4).value,  # Company
            ws_summary.cell(row=r, column=5).value,  # Vessel
            ws_summary.cell(row=r, column=6).value,  # Buyer
            ws_summary.cell(row=r, column=7).value   # End User
        ]

        # Get Status
        status = str(ws_summary.cell(row=r, column=BQ_col).value or "").strip().lower()

        if status == "plan":
            ws_b = ws_plan
            label = "PLAN"
        elif status in ("loading", "in progress", "completed"):
            ws_b = ws_comp
            label = "COMP"
        else:
            continue

        fill_hex = extract_fill_hex(cell_anq)

        ws_b.append([*key, val, fill_hex])

        print(f"[BACKUP] Row {r} → sheet {label}")
        print(f"         Key   : {key}")
        print(f"         Value : {val}")
        print(f"         Color : {fill_hex}\n")


def restore_fpg(ws_summary: Worksheet, plan_sheet_name: str, comp_sheet_name: str, start_row: int, end_row: int):
    """
    Restore FPG Actual Value (ANQ) with the color fill from FPG_plan dan FPG_comp sheet.
    """
    wb = ws_summary.parent

    ws_plan = wb[plan_sheet_name] if plan_sheet_name in wb.sheetnames else None
    ws_comp = wb[comp_sheet_name] if comp_sheet_name in wb.sheetnames else None

    if end_row is None:
        end_row = ws_summary.max_row

    try:
        ANQ_col = column_index_from_string("ANQ")
    except ValueError:
        raise ValueError("ERROR: 'ANQ' is not valid Excel Column Name")

    print(f"\n=== ♻ RESTORE FPG VALUE ({start_row} → {end_row}) ===")

    # Get key from summary
    def get_key_summary(row):
        return tuple(str(ws_summary.cell(row=row, column=c).value or "").strip() for c in range(3, 8))

    def get_key_from_backup(ws, row):
        return tuple(str(ws.cell(row=row, column=c).value or "").strip() for c in range(1, 6))

    # === Helper: Convert backup text → openpyxl fill ===
    def make_fill(color_code: str):
        """
        Convert saved fill color string from backup_fpg() 
        into PatternFill object for openpyxl.
        """
        if not color_code or color_code == "UNKNOWN":
            return PatternFill(fill_type=None)

        # 1. ARGB / RGB
        if len(color_code) in (6, 8) and all(c in "0123456789ABCDEFabcdef" for c in color_code):
            return PatternFill(start_color=color_code, end_color=color_code, fill_type="solid")

        # 2. THEME color: THEME_x_TINT_y
        if color_code.startswith("THEME_"):
            try:
                parts = color_code.split("_")
                theme_index = int(parts[1])
                tint_value = float(parts[3])

                color_obj = Color(theme=theme_index, tint=tint_value)
                return PatternFill(start_color=color_obj, end_color=color_obj, fill_type="solid")
            except:
                print(f"    (WARN) Cannot parse theme color: {color_code}")
                return PatternFill(fill_type=None)

        # 3. Indexed color
        if color_code.startswith("INDEXED_"):
            try:
                index = int(color_code.split("_")[1])
                color_obj = Color(indexed=index)
                return PatternFill(start_color=color_obj, end_color=color_obj, fill_type="solid")
            except:
                print(f"    (WARN) Cannot parse indexed color: {color_code}")
                return PatternFill(fill_type=None)

        # Unknown → skip
        return PatternFill(fill_type=None)

    # Restore function for every backup sheet
    def restore_from_backup(ws_backup, label):
        if not ws_backup:
            print(f"[RESTORE] Sheet {label} NOT FOUND !.")
            return

        print(f"\n--- Restore from sheet {label} ---")

        for r_b in range(2, ws_backup.max_row + 1):
            backup_key = get_key_from_backup(ws_backup, r_b)
            backup_val = ws_backup.cell(row=r_b, column=6).value
            backup_color = ws_backup.cell(row=r_b, column=7).value

            if backup_val is None:
                continue

            print(f"[CHECK] Backup row {r_b}")
            print(f"        Key   : {backup_key}")
            print(f"        Value : {backup_val}")
            print(f"        Color : {backup_color}")

            # Look for matching row in summary
            for r_sum in range(start_row, end_row + 1):
                if get_key_summary(r_sum) != backup_key:
                    continue

                print(f"   → MATCH in summary row {r_sum}, restore FPG Actual data!")

                cell = ws_summary.cell(row=r_sum, column=ANQ_col)
                cell.value = backup_val

                # Convert backup color into actual fill
                fill_obj = make_fill(backup_color)
                cell.fill = fill_obj

                print(f"   → Restored to summary row {r_sum}")
                print(f"     ✔ value={backup_val}, color={backup_color}\n")

    restore_from_backup(ws_plan, "FPG_plan")
    restore_from_backup(ws_comp, "FPG_comp")